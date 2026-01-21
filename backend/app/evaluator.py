# app/evaluator.py
from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, Any, List, Tuple, Optional

PASS_THRESHOLD = 75

ALIASES: Dict[str, List[str]] = {
    "abstract": ["abstract", "summary"],
    "introduction": ["introduction", "background", "overview"],
    "literature_review": [
        "literature review",
        "review of literature",
        "review of related literature",
        "related literature",
        "related work",
    ],
    "methods": [
        "methods",
        "methodology",
        "materials and methods",
        "experimental procedures",
        "research design",
        "research methodology",
    ],
    "results": ["results", "findings", "results and analysis", "results & analysis"],
    "discussion": ["discussion", "analysis", "interpretation", "discussion and analysis"],
    "conclusion": ["conclusion", "conclusions", "concluding remarks", "summary and conclusion"],
    "references": ["references", "reference", "bibliography", "works cited", "literature cited"],
}

FORMATS: Dict[str, Dict[str, Any]] = {
    "imrad": {
        "required": ["abstract", "introduction", "methods", "results", "discussion", "references"],
        "min_words": {"abstract": 140, "introduction": 260, "methods": 220, "results": 180, "discussion": 220, "references": 40},
        "weights": {
            "presence": {"abstract": 12, "introduction": 14, "methods": 16, "results": 16, "discussion": 16, "references": 10},
            "depth": {"abstract": 6, "introduction": 6, "methods": 6, "results": 6, "discussion": 6},
            "citations": 8,
            "references_entries": 4,
            "order": 4,
        },
    },
    "chapter_3": {
        "required": ["introduction", "literature_review", "methods", "references"],
        "min_words": {"introduction": 300, "literature_review": 300, "methods": 250, "references": 40},
        "weights": {
            "presence": {"introduction": 22, "literature_review": 22, "methods": 24, "references": 10},
            "depth": {"introduction": 6, "literature_review": 6, "methods": 6},
            "citations": 8,
            "references_entries": 4,
            "order": 4,
        },
    },
    "chapter_4": {
        "required": ["introduction", "literature_review", "methods", "results", "references"],
        "min_words": {"introduction": 300, "literature_review": 300, "methods": 250, "results": 200, "references": 40},
        "weights": {
            "presence": {"introduction": 18, "literature_review": 18, "methods": 20, "results": 20, "references": 10},
            "depth": {"introduction": 6, "literature_review": 6, "methods": 6, "results": 6},
            "citations": 8,
            "references_entries": 4,
            "order": 4,
        },
    },
    "chapter_5": {
        "required": ["introduction", "literature_review", "methods", "results", "discussion", "references"],
        "min_words": {"introduction": 300, "literature_review": 300, "methods": 250, "results": 200, "discussion": 220, "references": 40},
        "weights": {
            "presence": {"introduction": 16, "literature_review": 16, "methods": 18, "results": 18, "discussion": 18, "references": 10},
            "depth": {"introduction": 6, "literature_review": 6, "methods": 6, "results": 6, "discussion": 6},
            "citations": 8,
            "references_entries": 4,
            "order": 4,
        },
    },
}


def _file_kind(path: str) -> str:
    ext = (Path(path).suffix or "").lower().lstrip(".")
    return ext if ext in ("pdf", "docx") else "unknown"


def _normalize_heading(s: str) -> str:
    x = (s or "").strip()
    x = re.sub(r"^\s*\(?\s*\d+(\.\d+)*\s*\)?\s*", "", x)
    x = x.strip().strip(":").strip("-").strip()
    x = re.sub(r"\s+", " ", x)
    return x.lower()


def _heading_to_key(norm: str) -> str:
    n = norm or ""
    if ("results" in n) and ("discussion" in n):
        return "results_discussion"
    for k, vals in ALIASES.items():
        for a in vals:
            if n == a:
                return k
    for k, vals in ALIASES.items():
        for a in vals:
            if n.startswith(a) or (a in n and len(n) <= len(a) + 14):
                return k
    return ""


def _looks_like_heading(line: str) -> bool:
    s = (line or "").strip()
    if not s:
        return False
    if len(s) > 110:
        return False
    if re.search(r"[.]{3,}", s):
        return False
    if s.count(" ") > 14:
        return False
    letters = sum(1 for c in s if c.isalpha())
    if letters < 3:
        return False
    upper_letters = sum(1 for c in s if c.isalpha() and c.isupper())
    if letters >= 8 and (upper_letters / max(1, letters)) >= 0.75:
        return True
    if re.fullmatch(r"[\w\s&:/\-()]{3,110}", s) and s[:1].isupper():
        return True
    return False


def _word_count(text: str) -> int:
    toks = re.findall(r"[A-Za-z0-9]+(?:'[A-Za-z0-9]+)?", text or "")
    return len(toks)


def _symbol_ratio(text: str) -> float:
    if not text:
        return 1.0
    total = len(text)
    good = sum(1 for c in text if c.isalnum() or c.isspace())
    bad = max(0, total - good)
    return round(bad / max(1, total), 4)


def _citation_signals(text: str) -> int:
    t = text or ""
    patterns = [
        r"\([A-Z][A-Za-z\-]+,\s*\d{4}[a-z]?\)",
        r"\([A-Z][A-Za-z\-]+\s+et\s+al\.,\s*\d{4}[a-z]?\)",
        r"\b[A-Z][A-Za-z\-]+ et al\.\s*\(\d{4}[a-z]?\)",
        r"\[[0-9]{1,3}\]",
        r"\bdoi:\s*10\.\d{4,9}/[-._;()/:A-Za-z0-9]+\b",
        r"\b[A-Z][A-Za-z\-]+(?:\s+et\s+al\.)?\s*\(\s*\d{4}[a-z]?\s*\)",
    ]
    total = 0
    for p in patterns:
        total += len(re.findall(p, t))
    return total


def _read_docx_all(path: str) -> Tuple[str, List[Dict[str, Any]]]:
    try:
        from docx import Document
    except ModuleNotFoundError as e:
        raise RuntimeError("missing_dependency:python-docx") from e

    doc = Document(path)
    blocks: List[Dict[str, Any]] = []

    def add_block(text: str, style_name: str):
        t = (text or "").strip()
        if not t:
            return
        sn = (style_name or "").strip()
        is_heading = sn.lower().startswith("heading") or _looks_like_heading(t)
        level = 0
        m = re.search(r"heading\s+(\d+)", sn.lower())
        if m:
            level = int(m.group(1))
        blocks.append({"text": t, "is_heading": bool(is_heading), "level": level, "style": sn})

    for p in doc.paragraphs:
        add_block(p.text, p.style.name if p.style else "")

    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    add_block(p.text, p.style.name if p.style else "")

    raw = "\n".join(b["text"] for b in blocks).strip()
    return raw, blocks


def _read_pdf_text(path: str) -> str:
    txt = ""
    try:
        import pdfplumber  # type: ignore
        parts: List[str] = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                t = page.extract_text() or ""
                if t.strip():
                    parts.append(t)
        txt = "\n".join(parts).strip()
    except Exception:
        txt = ""

    if txt:
        return txt

    try:
        from PyPDF2 import PdfReader  # type: ignore
        reader = PdfReader(path)
        parts2: List[str] = []
        for p in reader.pages:
            t = p.extract_text() or ""
            if t.strip():
                parts2.append(t)
        return "\n".join(parts2).strip()
    except Exception:
        return ""


def _split_sections_docx(blocks: List[Dict[str, Any]]) -> Dict[str, str]:
    sections: Dict[str, List[str]] = {}
    cur = "preamble"
    sections[cur] = []
    for b in blocks:
        t = b["text"]
        if b.get("is_heading"):
            key = _heading_to_key(_normalize_heading(t))
            if key:
                cur = key
                sections.setdefault(cur, [])
                continue
        sections.setdefault(cur, []).append(t)
    return {k: "\n".join(v).strip() for k, v in sections.items()}


def _split_sections_text(raw: str) -> Dict[str, str]:
    lines = [ln.rstrip() for ln in (raw or "").splitlines()]
    cleaned: List[str] = []
    for ln in lines:
        s = ln.strip()
        if not s:
            continue
        cleaned.append(s)

    sections: Dict[str, List[str]] = {}
    cur = "preamble"
    sections[cur] = []

    for ln in cleaned:
        if _looks_like_heading(ln):
            key = _heading_to_key(_normalize_heading(ln))
            if key:
                cur = key
                sections.setdefault(cur, [])
                continue
        sections.setdefault(cur, []).append(ln)

    return {k: "\n".join(v).strip() for k, v in sections.items()}


def _preflight(raw: str, kind: str, units_count: int) -> Dict[str, Any]:
    wc = _word_count(raw)
    alpha_num = sum(1 for c in (raw or "") if c.isalnum())
    sym = _symbol_ratio(raw)

    verdict = "ok"
    confidence = "high"
    message = "Text looks readable."

    if wc < 60 or alpha_num < 200:
        verdict = "empty_or_nonsense"
        confidence = "low"
        message = "We could not extract enough readable text."
    elif kind == "pdf" and wc < 150:
        verdict = "scanned"
        confidence = "low"
        message = "PDF extraction looks weak. This may be scanned and needs OCR."
    elif wc < 220:
        verdict = "too_short"
        confidence = "medium"
        message = "Text extracted, but coverage looks low."
    elif sym > 0.33 and wc < 420:
        verdict = "empty_or_nonsense"
        confidence = "low"
        message = "Extracted text looks noisy or incomplete."

    return {
        "verdict": verdict,
        "confidence": confidence,
        "message": message,
        "kind": kind,
        "word_count": wc,
        "units_count": units_count,
        "symbol_ratio": sym,
    }


def _detect_best_format(sections: Dict[str, str]) -> str:
    def has(k: str) -> bool:
        return bool((sections.get(k) or "").strip())

    imrad_hits = 0
    if has("abstract"):
        imrad_hits += 1
    if has("introduction"):
        imrad_hits += 1
    if has("methods"):
        imrad_hits += 1
    if has("results") or has("results_discussion"):
        imrad_hits += 1
    if has("discussion") or has("results_discussion"):
        imrad_hits += 1
    if has("references"):
        imrad_hits += 1

    chap_hits = 0
    if has("literature_review"):
        chap_hits += 2
    if has("introduction"):
        chap_hits += 1
    if has("methods"):
        chap_hits += 1
    if has("results") or has("results_discussion"):
        chap_hits += 1
    if has("discussion") or has("results_discussion"):
        chap_hits += 1

    if imrad_hits >= 4 and imrad_hits >= chap_hits:
        return "imrad"
    if chap_hits >= 4:
        if has("discussion") or has("results_discussion"):
            return "chapter_5"
        if has("results") or has("results_discussion"):
            return "chapter_4"
        return "chapter_3"
    return "chapter_3"


def _level(score: int) -> str:
    if score >= 75:
        return "Good"
    if score >= 55:
        return "Needs work"
    return "Critical"


def _decision(score: int) -> str:
    return "Clear" if score >= PASS_THRESHOLD else "Revise"


def _compute_max_points(spec: Dict[str, Any]) -> int:
    w = spec["weights"]
    total = 0
    total += sum(int(v) for v in (w.get("presence") or {}).values())
    total += sum(int(v) for v in (w.get("depth") or {}).values())
    total += int(w.get("citations") or 0)
    total += int(w.get("references_entries") or 0)
    total += int(w.get("order") or 0)
    return max(1, total)


def evaluate_document(file_path: str, fmt: str = "") -> Dict[str, Any]:
    p = Path(file_path)
    if not p.exists():
        raise RuntimeError(f"file_not_found:{file_path}")
    if p.stat().st_size <= 0:
        raise RuntimeError(f"file_empty:{file_path}")

    kind = _file_kind(file_path)
    raw = ""
    blocks: Optional[List[Dict[str, Any]]] = None

    if kind == "docx":
        raw, blocks = _read_docx_all(file_path)
    elif kind == "pdf":
        raw = _read_pdf_text(file_path)
    else:
        raw = ""

    raw = (raw or "").strip()
    raw_preview = (raw[:520] + "…") if len(raw) > 520 else raw

    sections = _split_sections_docx(blocks) if blocks is not None else _split_sections_text(raw)
    units_count = len([k for k, v in sections.items() if k != "preamble" and (v or "").strip()])
    pre = _preflight(raw, kind, units_count)

    requested = (fmt or "").strip().lower()
    if requested in ("", "auto", "detect"):
        used_format = _detect_best_format(sections)
    else:
        used_format = requested if requested in FORMATS else _detect_best_format(sections)

    detected = _detect_best_format(sections)
    spec = FORMATS.get(used_format, FORMATS["chapter_3"])
    max_points = _compute_max_points(spec)

    def sec_text(key: str) -> str:
        if key == "results":
            return (sections.get("results") or sections.get("results_discussion") or "").strip()
        if key == "discussion":
            return (sections.get("discussion") or sections.get("results_discussion") or "").strip()
        return (sections.get(key) or "").strip()

    def sec_wc(key: str) -> int:
        return _word_count(sec_text(key))

    def add_check(
        lst: List[Dict[str, Any]],
        cid: str,
        name: str,
        category: str,
        status: str,
        why: str,
        evidence: str,
        points: int,
        maxp: int,
    ) -> None:
        lst.append(
            {
                "id": cid,
                "name": name,
                "category": category,
                "status": status,
                "why": why,
                "evidence": evidence,
                "points": int(points),
                "max_points": int(maxp),
            }
        )

    checks: List[Dict[str, Any]] = []
    blockers: List[Dict[str, Any]] = []
    strengths: List[str] = []
    next_actions: List[str] = []

    cite_sig = _citation_signals(raw)

    if pre["verdict"] in ("empty_or_nonsense", "scanned"):
        summary = {
            "score": 0,
            "level": "Critical",
            "decision": "Revise",
            "threshold": PASS_THRESHOLD,
            "format": used_format,
            "format_detected": detected,
            "missing_required": [],
            "word_count": pre["word_count"],
            "citation_signals": cite_sig,
            "preflight": pre,
            "debug": {"file_kind": kind, "units_count": units_count, "raw_len": len(raw), "raw_preview": raw_preview},
        }
        add_check(checks, "preflight", "Text readability", "Preflight", "FAIL", pre["message"], f"word_count={pre['word_count']} symbol_ratio={pre['symbol_ratio']}", 0, 0)
        blockers.append({"title": "Unreadable text", "reason": pre["message"], "evidence": raw_preview or "No text extracted."})
        next_actions.append("Upload a DOCX exported from Word/Google Docs (best accuracy).")
        if kind == "pdf":
            next_actions.append("If the PDF is scanned, OCR it first or convert to DOCX.")
        next_actions.append("Use real headings and normal paragraphs (avoid textboxes/shapes).")
        return {"summary": summary, "checks": checks, "blockers": blockers, "strengths": strengths, "connections": {}, "fix_plan": {"next_actions": next_actions}}

    missing_required: List[str] = []

    presence_w = spec["weights"]["presence"]
    for key in spec["required"]:
        label = key.replace("_", " ").title()
        present = bool(sec_text(key))
        maxp = int(presence_w.get(key, 0))
        if present:
            add_check(checks, f"presence_{key}", f"{label} section", "Structure", "PASS", "Section detected.", f"{sec_wc(key)} words", maxp, maxp)
        else:
            missing_required.append(label)
            add_check(checks, f"presence_{key}", f"{label} section", "Structure", "FAIL", "Required section not detected.", f"Expected heading like: {label}", 0, maxp)

    depth_w = spec["weights"].get("depth") or {}
    min_words = spec.get("min_words") or {}
    for key, maxp in depth_w.items():
        label = key.replace("_", " ").title()
        text = sec_text(key)
        need = int(min_words.get(key, 180))
        if not text:
            add_check(checks, f"depth_{key}", f"{label} depth", "Content", "FAIL", "Section text not found.", "No text captured for this section.", 0, int(maxp))
            continue
        w = _word_count(text)
        if w >= need:
            add_check(checks, f"depth_{key}", f"{label} depth", "Content", "PASS", "Content depth looks developed.", f"{w} words (min ~{need})", int(maxp), int(maxp))
        elif w >= int(need * 0.6):
            add_check(checks, f"depth_{key}", f"{label} depth", "Content", "WEAK", "Section exists but needs more depth.", f"{w} words (min ~{need})", max(0, int(maxp) - 2), int(maxp))
        else:
            add_check(checks, f"depth_{key}", f"{label} depth", "Content", "FAIL", "Section is too short to be reliable.", f"{w} words (min ~{need})", 0, int(maxp))

    order_pts = int(spec["weights"].get("order") or 0)
    order_score = 0
    order_msg = "Section order looks reasonable."
    if used_format == "imrad":
        seq = ["abstract", "introduction", "methods", "results", "discussion", "references"]
    elif used_format == "chapter_3":
        seq = ["introduction", "literature_review", "methods", "references"]
    elif used_format == "chapter_4":
        seq = ["introduction", "literature_review", "methods", "results", "references"]
    else:
        seq = ["introduction", "literature_review", "methods", "results", "discussion", "references"]

    present_seq = [k for k in seq if sec_text(k)]
    if len(present_seq) >= 3:
        order_score = order_pts
    else:
        order_msg = "Not enough sections detected to validate order."
        order_score = max(0, order_pts - 3)

    add_check(
        checks,
        "order_check",
        "Section order",
        "Structure",
        "PASS" if order_score == order_pts else "WEAK",
        order_msg,
        f"detected_in_order={', '.join(present_seq) if present_seq else '—'}",
        int(order_score),
        int(order_pts),
    )

    cite_max = int(spec["weights"].get("citations") or 0)
    if cite_sig >= 10:
        add_check(checks, "citations_signals", "In-text citations", "Citations", "PASS", "Strong citation signals detected.", f"signals={cite_sig}", cite_max, cite_max)
    elif cite_sig >= 3:
        add_check(checks, "citations_signals", "In-text citations", "Citations", "WEAK", "Some citations detected, but evidence looks light.", f"signals={cite_sig}", max(0, cite_max - 2), cite_max)
    else:
        add_check(checks, "citations_signals", "In-text citations", "Citations", "FAIL", "No reliable in-text citation patterns detected.", f"signals={cite_sig}", 0, cite_max)

    refs_max = int(spec["weights"].get("references_entries") or 0)
    ref_text = sec_text("references")
    entries = 0
    if ref_text:
        lines = [ln.strip() for ln in ref_text.splitlines() if ln.strip()]
        entries = len([ln for ln in lines if len(ln) >= 12])
    else:
        tail = "\n".join((raw or "").splitlines()[-80:])
        if re.search(r"\breferences\b|\bbibliography\b|\bworks cited\b", tail, flags=re.I):
            entries = max(entries, 2)

    if entries >= 6:
        add_check(checks, "references_present", "References present", "Citations", "PASS", "References detected with multiple entries.", f"entries≈{entries}", refs_max, refs_max)
    elif entries >= 2:
        add_check(checks, "references_present", "References present", "Citations", "WEAK", "References detected but entries look few.", f"entries≈{entries}", max(0, refs_max - 1), refs_max)
    else:
        add_check(checks, "references_present", "References present", "Citations", "FAIL", "References section/entries not reliably detected.", "Add a clear heading: References / Bibliography", 0, refs_max)

    earned = sum(int(c.get("points") or 0) for c in checks)
    score = int(round((earned / max_points) * 100))
    score = max(0, min(100, score))

    for c in checks:
        if c["status"] == "FAIL":
            blockers.append({"title": c["name"], "reason": c["why"], "evidence": c["evidence"]})
        elif c["status"] == "PASS":
            strengths.append(c["name"])

    blockers = blockers[:6]
    strengths = strengths[:10]

    for b in blockers[:4]:
        next_actions.append(f"Fix: {b['title']}")

    if requested not in ("", "auto", "detect") and requested != used_format:
        next_actions.insert(0, f"Format mismatch: selected “{requested}” but content fits “{used_format}” better.")

    if not next_actions:
        next_actions.append("Looks good. Re-upload after small edits to confirm stability.")

    summary = {
        "score": score,
        "level": _level(score),
        "decision": _decision(score),
        "threshold": PASS_THRESHOLD,
        "format": used_format,
        "format_detected": detected,
        "missing_required": missing_required,
        "word_count": pre["word_count"],
        "citation_signals": cite_sig,
        "preflight": pre,
        "debug": {
            "file_kind": kind,
            "units_count": units_count,
            "raw_len": len(raw),
            "raw_preview": raw_preview,
            "max_points": max_points,
            "earned_points": earned,
        },
    }

    return {
        "summary": summary,
        "checks": checks,
        "blockers": blockers,
        "strengths": strengths,
        "connections": {},
        "fix_plan": {"next_actions": next_actions},
    }
